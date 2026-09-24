import csv
import os
from openpyxl import load_workbook


# ============================================================
# FILES
# ============================================================

INPUT_FILE = "data/us_1M.csv"
KIR_FILE = "data/Kirs.xlsx"

base, ext = os.path.splitext(INPUT_FILE)
OUTPUT_FILE = base + "_kir" + ext


# ============================================================
# LOAD KIR REFERENCE
# ============================================================

print("=" * 70)
print("LOADING KIR REFERENCE")
print("=" * 70)

wb = load_workbook(KIR_FILE, read_only=True, data_only=True)


def load_group(sheet_name, group_name):
    """
    Read a sheet from Kirs.xlsx.

    Column A = HLA allele
    Column C = KIR ligand group

    Return alleles whose column C equals group_name.
    """

    ws = wb[sheet_name]
    alleles = set()

    for row in ws.iter_rows(values_only=True):
        if len(row) < 3:
            continue

        allele = row[0]
        group = row[2]

        if allele is None or group is None:
            continue

        allele = str(allele).strip()
        group = str(group).strip()

        if group == group_name:
            alleles.add(allele)

    return alleles


# Bw4 can be associated with HLA-A and HLA-B
bw4_A = load_group("Aw", "Bw4")
bw4_B = load_group("Bw", "Bw4")

# C1 is determined from HLA-C
c1_C = load_group("Cw", "C1")


print(f"A alleles with Bw4 : {len(bw4_A):,}")
print(f"B alleles with Bw4 : {len(bw4_B):,}")
print(f"C alleles with C1  : {len(c1_C):,}")


# ============================================================
# PARSE GENOTYPE
# ============================================================

def parse_genotype(geno):
    """
    Example genotype:

    A*11:01+A*26:01^
    B*27:05+B*38:01^
    C*02:02+C*12:03^
    DQB1*03:03+DQB1*05:02^
    DRB1*09:01+DRB1*16:01

    Returns:
        A = [A1, A2]
        B = [B1, B2]
        C = [C1, C2]
    """

    loci = geno.split("^")

    if len(loci) < 3:
        raise ValueError(f"Invalid genotype: {geno}")

    A = [x.strip() for x in loci[0].split("+")]
    B = [x.strip() for x in loci[1].split("+")]
    C = [x.strip() for x in loci[2].split("+")]

    return A, B, C


# ============================================================
# COUNTERS
# ============================================================

n_rows = 0

bw4_distribution = {
    0: 0,
    1: 0,
    2: 0,
    3: 0,
    4: 0,
}

c1_distribution = {
    0: 0,
    1: 0,
    2: 0,
}

# Useful for checking the final combined KIR values
kir_distribution = {}


# ============================================================
# PROCESS DATA
# ============================================================

print()
print("=" * 70)
print("PROCESSING DATA")
print("=" * 70)

print(f"Input : {INPUT_FILE}")
print(f"Output: {OUTPUT_FILE}")
print()


with open(INPUT_FILE, "r", newline="") as fin, \
     open(OUTPUT_FILE, "w", newline="") as fout:

    reader = csv.reader(fin)
    writer = csv.writer(fout)

    for row_number, row in enumerate(reader, start=1):

        if not row:
            continue

        if len(row) < 2:
            raise ValueError(
                f"Row {row_number}: expected at least 2 columns, got {len(row)}"
            )

        # ----------------------------------------------------
        # Original data
        # ----------------------------------------------------

        person_id = row[0].strip()
        geno = row[1].strip()

        # ----------------------------------------------------
        # Extract A, B, C alleles
        # ----------------------------------------------------

        A, B, C = parse_genotype(geno)

        if len(A) != 2:
            raise ValueError(
                f"Row {row_number}: expected 2 A alleles, got {A}"
            )

        if len(B) != 2:
            raise ValueError(
                f"Row {row_number}: expected 2 B alleles, got {B}"
            )

        if len(C) != 2:
            raise ValueError(
                f"Row {row_number}: expected 2 C alleles, got {C}"
            )

        # ----------------------------------------------------
        # Bw4
        #
        # Check all four:
        # A1, A2, B1, B2
        #
        # Result can be 0..4
        # ----------------------------------------------------

        bw4_count = (
            sum(allele in bw4_A for allele in A)
            +
            sum(allele in bw4_B for allele in B)
        )

        # ----------------------------------------------------
        # C1
        #
        # Check:
        # C1, C2
        #
        # Result can be 0..2
        # ----------------------------------------------------

        c1_count = sum(
            allele in c1_C
            for allele in C
        )

        # ----------------------------------------------------
        # Scenario1 KIR format
        #
        # Examples:
        # 0;0
        # 1;0
        # 1;1
        # 2;1
        # 2;2
        # ----------------------------------------------------

        kir = f"{bw4_count};{c1_count}"

        # ----------------------------------------------------
        # Write EXACTLY:
        #
        # id,geno,kir
        #
        # We intentionally DO NOT copy frequency/weight/etc.
        # from the original file.
        # ----------------------------------------------------

        writer.writerow([
            person_id,
            geno,
            kir,
        ])

        # ----------------------------------------------------
        # Statistics
        # ----------------------------------------------------

        bw4_distribution[bw4_count] += 1
        c1_distribution[c1_count] += 1

        kir_distribution[kir] = kir_distribution.get(kir, 0) + 1

        n_rows += 1

        if n_rows % 100000 == 0:
            print(f"Processed {n_rows:,} rows")


# ============================================================
# DONE
# ============================================================

wb.close()

print()
print("=" * 70)
print("DONE")
print("=" * 70)

print(f"Rows processed: {n_rows:,}")
print(f"Output file   : {OUTPUT_FILE}")


# ============================================================
# DISTRIBUTIONS
# ============================================================

print()
print("Bw4 count distribution")
print("-" * 30)

for value in sorted(bw4_distribution):
    count = bw4_distribution[value]
    percent = 100.0 * count / n_rows if n_rows else 0

    print(
        f"Bw4={value}: "
        f"{count:,} "
        f"({percent:.2f}%)"
    )


print()
print("C1 count distribution")
print("-" * 30)

for value in sorted(c1_distribution):
    count = c1_distribution[value]
    percent = 100.0 * count / n_rows if n_rows else 0

    print(
        f"C1={value}: "
        f"{count:,} "
        f"({percent:.2f}%)"
    )


print()
print("Combined KIR distribution")
print("-" * 30)

for kir in sorted(kir_distribution):
    count = kir_distribution[kir]
    percent = 100.0 * count / n_rows if n_rows else 0

    print(
        f"{kir}: "
        f"{count:,} "
        f"({percent:.2f}%)"
    )